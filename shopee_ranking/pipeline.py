#!/usr/bin/env python3
"""Shopee Indonesia 売れ筋ランキング DB パイプライン.

公式 Shopee Affiliate ランキング（Produk Trending = 7D / Produk Terlaris = 30D）を
認証済みセッションで取得した CSV を inbox/ から取り込み、追記専用の履歴 CSV と
5 シート構成の xlsx（Google Sheets 取り込み用）を更新する。

取得できなかった実行も --fail-reason で Run_Log に記録する。架空データは生成しない。

  python3 pipeline.py ingest                 # inbox/*.csv を取り込み
  python3 pipeline.py ingest --fail-reason "Affiliate 未ログイン"   # 取得失敗を記録
  python3 pipeline.py build                  # 派生シートと xlsx を再生成のみ
"""
import argparse
import csv
import re
import shutil
import sys
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

WIB = ZoneInfo("Asia/Jakarta")
ROOT = Path(__file__).resolve().parent
DATA = ROOT / "data"
INBOX = ROOT / "inbox"
PROCESSED = INBOX / "processed"
WORKBOOK = DATA / "Shopee_ID_Ranking_DB.xlsx"

TOP_N = 30
PERIODS = {"7D": "Ranking_7D", "30D": "Ranking_30D"}
OFFICIAL_SOURCES = {
    "affiliate_trending": ("7D", "Shopee Affiliate - Produk Trending（販売増加）"),
    "affiliate_terlaris": ("30D", "Shopee Affiliate - Produk Terlaris（販売数量）"),
    "mcn": (None, "Shopee MCN ランキング"),
}
# 公開売れ筋ページは公式ランキングとして扱わず別管理する
REFERENCE_SOURCES = {"public_bestseller"}

RANKING_COLS = [
    "取得日時_WIB", "取得日", "対象期間", "公式順位", "商品名_原文", "商品名_日本語",
    "カテゴリー", "販売価格_Rp", "販売数量", "販売数量_対象範囲", "推定売上高_Rp",
    "店舗名", "商品URL", "商品ID", "店舗ID", "レビュー評価", "レビュー件数",
    "前回順位", "順位変動", "新規ランクイン", "輸出適性", "輸出適性_根拠",
    "取得元", "データ取得元URL", "取得条件", "取得状況_注意事項", "run_id",
]
RUN_LOG_COLS = [
    "run_id", "実行日時_WIB", "対象期間", "取得元", "データ取得元URL", "取得条件",
    "期待件数", "取得件数", "欠損値", "エラー", "実行結果", "備考",
]
TREND_COLS = [
    "対象期間", "最新取得日", "商品ID", "商品名_原文", "商品名_日本語", "最新順位",
    "前回順位", "順位変動", "新規ランクイン", "連続ランクイン日数", "通算ランクイン回数",
    "初登場日", "最新価格_Rp", "前回価格_Rp", "価格変動率_%", "両ランキング登場", "注目理由",
]
REFERENCE_COLS = ["取得日時_WIB", "取得日", "商品名_原文", "販売価格_Rp", "販売数量",
                  "商品URL", "商品ID", "店舗ID", "データ取得元URL", "注意事項"]

# inbox CSV で受け付ける列（英語名 → DB 列）
INBOX_MAP = {
    "rank": "公式順位", "product_name": "商品名_原文", "product_name_ja": "商品名_日本語",
    "category": "カテゴリー", "price_rp": "販売価格_Rp", "sales_qty": "販売数量",
    "sales_qty_scope": "販売数量_対象範囲", "shop_name": "店舗名", "product_url": "商品URL",
    "item_id": "商品ID", "shop_id": "店舗ID", "rating": "レビュー評価",
    "review_count": "レビュー件数", "export_potential": "輸出適性",
    "export_reason": "輸出適性_根拠", "notes": "取得状況_注意事項",
}
REQUIRED_FIELDS = ["商品名_原文", "販売価格_Rp", "店舗名", "商品URL", "商品ID", "店舗ID"]


def now_wib():
    return datetime.now(WIB).replace(microsecond=0)


def read_csv(path):
    if not path.exists():
        return []
    with path.open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def write_csv(path, cols, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=cols, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)


def append_csv(path, cols, rows):
    write_csv(path, cols, read_csv(path) + rows)


def parse_ids(url):
    """shopee.co.id/product/{shop}/{item} または ...-i.{shop}.{item} から ID を抽出."""
    m = re.search(r"/product/(\d+)/(\d+)", url or "") or re.search(r"-i\.(\d+)\.(\d+)", url or "")
    return (m.group(1), m.group(2)) if m else ("", "")


def to_number(text):
    """'Rp12.500' → 12500。範囲表記や不明値は None（推定しない）."""
    s = str(text or "").strip()
    if not s or "-" in s.replace("Rp", "").strip("- ") or "~" in s:
        return None
    digits = re.sub(r"[^\d]", "", s)
    if not digits or re.search(r"(rb|jt|k|\+)", s, re.I):
        return None
    return int(digits)


def item_key(row):
    return row.get("商品ID") or row.get("商品URL") or row.get("商品名_原文")


def read_inbox_file(path):
    """inbox CSV: 先頭の '#key: value' 行がメタデータ、その後がデータ."""
    meta, lines = {}, []
    with path.open(encoding="utf-8-sig") as f:
        for line in f:
            if line.startswith("#"):
                k, _, v = line[1:].partition(":")
                meta[k.strip()] = v.strip()
            elif line.strip():
                lines.append(line)
    return meta, list(csv.DictReader(lines))


def normalize(meta, raw_rows, run_id):
    source = meta.get("source", "")
    if source in REFERENCE_SOURCES:
        period = "REF"
    elif source in OFFICIAL_SOURCES:
        period = meta.get("period") or OFFICIAL_SOURCES[source][0]
    else:
        raise ValueError(f"未対応の source: {source!r}（{', '.join([*OFFICIAL_SOURCES, *REFERENCE_SOURCES])}）")
    if period not in (*PERIODS, "REF"):
        raise ValueError(f"period は 7D / 30D のみ: {period!r}")
    captured = meta.get("captured_at_wib") or now_wib().isoformat()
    captured_dt = datetime.fromisoformat(captured)
    if captured_dt.tzinfo is None:
        captured_dt = captured_dt.replace(tzinfo=WIB)
    captured_dt = captured_dt.astimezone(WIB)

    rows = []
    for raw in raw_rows:
        r = {db: (raw.get(src) or "").strip() for src, db in INBOX_MAP.items()}
        shop_id, item_id = parse_ids(r["商品URL"])
        r["店舗ID"] = r["店舗ID"] or shop_id
        r["商品ID"] = r["商品ID"] or item_id
        r.update({
            "取得日時_WIB": captured_dt.isoformat(), "取得日": captured_dt.date().isoformat(),
            "対象期間": period, "取得元": OFFICIAL_SOURCES.get(source, (None, source))[1],
            "データ取得元URL": meta.get("source_url", ""), "取得条件": meta.get("conditions", ""),
            "run_id": run_id,
        })
        # 売上高は販売数量が対象期間の値だと明示されている場合のみ計算
        price, qty = to_number(r["販売価格_Rp"]), to_number(r["販売数量"])
        if period in PERIODS and r["販売数量_対象範囲"].upper() == period and price and qty:
            r["推定売上高_Rp"] = str(price * qty)
        else:
            r["推定売上高_Rp"] = ""
            if r["販売数量"]:
                note = "売上高未計算（販売数量の対象期間が未確認、または価格が範囲表示）"
                r["取得状況_注意事項"] = "; ".join(filter(None, [r["取得状況_注意事項"], note]))
        r["輸出適性"] = r["輸出適性"] or "未評価"
        rows.append(r)
    return period, source, meta, rows


def validate(rows):
    errors, missing = [], defaultdict(int)
    ranks = []
    for r in rows:
        try:
            ranks.append(int(r["公式順位"]))
        except ValueError:
            errors.append(f"順位が数値でない: {r['商品名_原文'][:30]!r}")
        for f in REQUIRED_FIELDS:
            if not r.get(f):
                missing[f] += 1
    if len(ranks) != len(set(ranks)):
        errors.append("順位が重複")
    if any(not 1 <= k <= TOP_N for k in ranks):
        errors.append(f"1〜{TOP_N} 位の範囲外の順位あり")
    return errors, dict(missing)


def ingest(fail_reason=None):
    run_at = now_wib()
    run_id = run_at.strftime("%Y%m%d-%H%M%S")
    logs = []

    if fail_reason:
        for period, sheet in PERIODS.items():
            src = "affiliate_trending" if period == "7D" else "affiliate_terlaris"
            logs.append({
                "run_id": run_id, "実行日時_WIB": run_at.isoformat(), "対象期間": period,
                "取得元": OFFICIAL_SOURCES[src][1], "データ取得元URL": "https://affiliate.shopee.co.id/",
                "取得条件": "全カテゴリー / 全価格帯 / チャネル: すべて", "期待件数": TOP_N,
                "取得件数": 0, "欠損値": "", "エラー": fail_reason, "実行結果": "FAILED",
                "備考": "架空データでの補完なし",
            })
        append_csv(DATA / "Run_Log.csv", RUN_LOG_COLS, logs)
        build()
        return logs

    files = sorted(INBOX.glob("*.csv"))
    if not files:
        print("inbox に CSV がありません。取得失敗として記録するには --fail-reason を指定してください。")
        return []

    for path in files:
        if path.name.startswith("TEMPLATE"):
            continue
        meta, raw = read_inbox_file(path)
        try:
            period, source, meta, rows = normalize(meta, raw, run_id)
        except ValueError as e:
            logs.append({"run_id": run_id, "実行日時_WIB": run_at.isoformat(), "取得元": path.name,
                         "期待件数": TOP_N, "取得件数": 0, "エラー": str(e), "実行結果": "FAILED"})
            continue

        if period == "REF":
            append_csv(DATA / "Reference_Public.csv", REFERENCE_COLS, rows)
            logs.append({"run_id": run_id, "実行日時_WIB": run_at.isoformat(), "対象期間": "参考",
                         "取得元": source, "取得件数": len(rows), "実行結果": "REFERENCE",
                         "備考": "公開売れ筋ページ。公式ランキングとは別管理"})
            shutil.move(path, PROCESSED / path.name)
            continue

        errors, missing = validate(rows)
        hist_path = DATA / f"{PERIODS[period]}.csv"
        history = read_csv(hist_path)
        seen = {(h["取得日"], h["対象期間"], item_key(h)) for h in history}
        fresh = [r for r in rows if (r["取得日"], period, item_key(r)) not in seen]
        dup = len(rows) - len(fresh)
        if errors:
            status = "FAILED"
        elif not fresh and dup:
            status = "DUPLICATE"
        elif len(fresh) < TOP_N:
            status = "PARTIAL"
        else:
            status = "SUCCESS"
        if not errors:
            append_csv(hist_path, RANKING_COLS, fresh)
        notes = []
        if dup:
            notes.append(f"同日重複 {dup} 件をスキップ")
        if len(rows) < TOP_N:
            notes.append(f"取得 {len(rows)}/{TOP_N} 件。不足理由: {meta.get('shortfall_reason', '未記入')}")
        logs.append({
            "run_id": run_id, "実行日時_WIB": run_at.isoformat(), "対象期間": period,
            "取得元": OFFICIAL_SOURCES[source][1], "データ取得元URL": meta.get("source_url", ""),
            "取得条件": meta.get("conditions", ""), "期待件数": TOP_N, "取得件数": len(fresh),
            "欠損値": ", ".join(f"{k}:{v}" for k, v in missing.items()), "エラー": "; ".join(errors),
            "実行結果": status, "備考": "; ".join(notes),
        })
        if not errors:
            shutil.move(path, PROCESSED / path.name)

    append_csv(DATA / "Run_Log.csv", RUN_LOG_COLS, logs)
    build()
    return logs


def by_date(history):
    days = defaultdict(dict)
    for r in history:
        days[r["取得日"]][item_key(r)] = r
    return days


def annotate(history):
    """前回順位・順位変動・新規ランクインを、直前の取得日との比較で計算し直す."""
    days = by_date(history)
    order = sorted(days)
    for i, d in enumerate(order):
        prev = days[order[i - 1]] if i else None
        for k, r in days[d].items():
            p = prev.get(k) if prev else None
            r["前回順位"] = p["公式順位"] if p else ""
            r["順位変動"] = str(int(p["公式順位"]) - int(r["公式順位"])) if p else ""
            r["新規ランクイン"] = "" if prev is None else ("いいえ" if p else "はい")
    return days, order


def build():
    DATA.mkdir(parents=True, exist_ok=True)
    trend, latest = [], {}
    for period, sheet in PERIODS.items():
        path = DATA / f"{sheet}.csv"
        history = read_csv(path)
        days, order = annotate(history)
        write_csv(path, RANKING_COLS, history)
        if not order:
            latest[period] = []
            continue
        cur = days[order[-1]]
        prev = days[order[-2]] if len(order) > 1 else {}
        latest[period] = sorted(cur.values(), key=lambda r: int(r["公式順位"]))
        for k, r in cur.items():
            streak = 0
            for d in reversed(order):
                if k not in days[d]:
                    break
                streak += 1
            appearances = [d for d in order if k in days[d]]
            p = prev.get(k)
            price, pprice = to_number(r["販売価格_Rp"]), to_number(p["販売価格_Rp"]) if p else None
            trend.append({
                "対象期間": period, "最新取得日": order[-1], "商品ID": r["商品ID"],
                "商品名_原文": r["商品名_原文"], "商品名_日本語": r["商品名_日本語"],
                "最新順位": r["公式順位"], "前回順位": r["前回順位"], "順位変動": r["順位変動"],
                "新規ランクイン": r["新規ランクイン"], "連続ランクイン日数": streak,
                "通算ランクイン回数": len(appearances), "初登場日": appearances[0],
                "最新価格_Rp": r["販売価格_Rp"], "前回価格_Rp": p["販売価格_Rp"] if p else "",
                "価格変動率_%": f"{(price - pprice) / pprice * 100:.1f}" if price and pprice else "",
            })
    keys = {p: {item_key(r) for r in rows} for p, rows in latest.items()}
    both = keys.get("7D", set()) & keys.get("30D", set())
    for t in trend:
        t["両ランキング登場"] = "はい" if t["商品ID"] in both or t["商品名_原文"] in both else "いいえ"
        reasons = []
        if t["新規ランクイン"] == "はい":
            reasons.append("新規")
        if t["順位変動"] and int(t["順位変動"]) >= 5:
            reasons.append(f"{t['順位変動']}位上昇")
        if t["両ランキング登場"] == "はい":
            reasons.append("7D・30D 両方")
        t["注目理由"] = " / ".join(reasons)
    write_csv(DATA / "Trend_Analysis.csv", TREND_COLS, trend)

    side_cols = ["順位"] + [f"{p}_{c}" for p in PERIODS for c in
                            ("商品名_原文", "商品名_日本語", "販売価格_Rp", "店舗名", "順位変動", "新規ランクイン", "商品URL")]
    side = []
    for rank in range(1, TOP_N + 1):
        row = {"順位": rank}
        for p in PERIODS:
            r = next((x for x in latest.get(p, []) if int(x["公式順位"]) == rank), None)
            for c in ("商品名_原文", "商品名_日本語", "販売価格_Rp", "店舗名", "順位変動", "新規ランクイン", "商品URL"):
                row[f"{p}_{c}"] = r[c] if r else ("（未取得）" if c == "商品名_原文" else "")
        side.append(row)
    write_csv(DATA / "Latest_Top30.csv", side_cols, side)
    write_workbook(side_cols)


def write_workbook(side_cols):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        print("openpyxl 未インストールのため xlsx は省略（CSV は更新済み）", file=sys.stderr)
        return
    wb = Workbook()
    wb.remove(wb.active)
    sheets = [("Ranking_7D", RANKING_COLS), ("Ranking_30D", RANKING_COLS),
              ("Latest_Top30", side_cols), ("Trend_Analysis", TREND_COLS), ("Run_Log", RUN_LOG_COLS)]
    for name, cols in sheets:
        ws = wb.create_sheet(name)
        ws.append(cols)
        for r in read_csv(DATA / f"{name}.csv"):
            ws.append([r.get(c, "") for c in cols])
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="EE4D2D")
        ws.freeze_panes = "A2"
    wb.save(WORKBOOK)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("command", choices=["ingest", "build"])
    ap.add_argument("--fail-reason", help="データを取得できなかった理由（Run_Log に FAILED で記録）")
    args = ap.parse_args()
    PROCESSED.mkdir(parents=True, exist_ok=True)
    if args.command == "build":
        build()
    else:
        for log in ingest(args.fail_reason):
            print(f"[{log.get('実行結果')}] {log.get('対象期間', '')} {log.get('取得件数')}件 {log.get('エラー', '')}")


if __name__ == "__main__":
    main()
